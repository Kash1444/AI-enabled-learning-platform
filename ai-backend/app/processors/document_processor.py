"""
Document processing: text extraction, cleaning and chunking.

Supported formats: PDF (PyMuPDF/fitz), DOCX (python-docx), PPTX
(python-pptx), TXT. Adding a new format later only means adding one more
`extract_*` function and one `elif` branch in `extract_text`.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import List

from app.core.config import get_settings

settings = get_settings()


def extract_pdf(path: Path) -> str:
    # PyMuPDF renamed its module from `fitz` to `pymupdf`; `fitz` still works
    # but warns on import and is slated for removal.
    try:
        import pymupdf
    except ImportError:  # pragma: no cover - PyMuPDF < 1.24
        import fitz as pymupdf

    text_parts: List[str] = []
    with pymupdf.open(path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def extract_docx(path: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts: List[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation  # python-pptx

    prs = Presentation(str(path))
    parts: List[str] = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        slide_text: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    run_text = "".join(run.text for run in paragraph.runs)
                    if run_text.strip():
                        slide_text.append(run_text)
        if slide_text:
            parts.append(f"[Slide {slide_number}] " + " ".join(slide_text))
    return "\n".join(parts)


def extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_text(path: Path, file_type: str) -> str:
    """Dispatch to the right extractor based on file extension."""
    ext = file_type.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".pptx":
        return extract_pptx(path)
    if ext == ".txt":
        return extract_txt(path)
    raise ValueError(f"Unsupported file type for extraction: {ext}")


def clean_text(raw_text: str) -> str:
    """Normalize whitespace and drop empty/near-empty lines."""
    text = raw_text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    lines = [ln.strip() for ln in text.split("\n")]
    lines = [ln for ln in lines if ln]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(
    text: str,
    chunk_size: int | None = None,
    overlap: int | None = None,
) -> List[str]:
    """
    Split text into overlapping character-based chunks.

    A character-based sliding window is simple, dependency-free and works
    well enough for MCQ generation / RAG retrieval at hackathon scale. Word
    boundaries are respected where possible to avoid cutting mid-word.
    """
    chunk_size = chunk_size or settings.RAG_CHUNK_SIZE
    overlap = overlap or settings.RAG_CHUNK_OVERLAP
    if chunk_size <= overlap:
        raise ValueError("chunk_size must be greater than overlap")

    if not text:
        return []

    chunks: List[str] = []
    start = 0
    length = len(text)
    while start < length:
        end = min(start + chunk_size, length)
        # try to end on a sentence/word boundary
        if end < length:
            boundary = text.rfind(" ", start, end)
            if boundary > start:
                end = boundary
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= length:
            break

        next_start = max(end - overlap, start + 1)
        # Stepping back by `overlap` characters usually lands in the middle of
        # a word, which would make the next chunk open with a fragment like
        # "ariance relative to...". Those fragments then surface verbatim as
        # MCQ options and RAG citations, so snap forward to the next whole
        # word. `next_start > start` already holds, so progress is guaranteed.
        if next_start < length and not text[next_start - 1].isspace():
            probe = next_start
            while probe < length and not text[probe].isspace():
                probe += 1
            while probe < length and text[probe].isspace():
                probe += 1
            if probe < length:
                next_start = probe
        start = next_start
    return chunks
