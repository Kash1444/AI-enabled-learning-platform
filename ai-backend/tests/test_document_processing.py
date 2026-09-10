from pathlib import Path

from app.processors.document_processor import chunk_text, clean_text, extract_text


def test_clean_text_normalizes_whitespace():
    raw = "Line one\r\n\r\n\r\n   Line   two  \n\n\nLine three\r"
    cleaned = clean_text(raw)
    assert "\r" not in cleaned
    assert "\n\n\n" not in cleaned
    assert "Line   two" not in cleaned  # internal multi-space collapsed


def test_chunk_text_respects_size_and_overlap():
    text = "word " * 1000
    chunks = chunk_text(text, chunk_size=200, overlap=50)
    assert len(chunks) > 1
    for c in chunks:
        assert len(c) <= 200 + 1  # small tolerance for boundary trimming


def test_chunk_text_empty_input_returns_empty_list():
    assert chunk_text("") == []


def test_chunk_text_rejects_overlap_gte_chunk_size():
    import pytest

    with pytest.raises(ValueError):
        chunk_text("some text", chunk_size=10, overlap=10)


def test_extract_txt(tmp_path: Path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text("Hello world.\nThis is a sample statistics training document.")
    text = extract_text(file_path, ".txt")
    assert "Hello world." in text
    assert "statistics" in text


def test_extract_docx(tmp_path: Path):
    import docx

    file_path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_paragraph("Sampling methodology overview.")
    document.add_paragraph("Stratified sampling reduces variance.")
    document.save(str(file_path))

    text = extract_text(file_path, ".docx")
    assert "Sampling methodology overview." in text
    assert "Stratified sampling reduces variance." in text


def test_extract_pptx(tmp_path: Path):
    from pptx import Presentation

    file_path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Survey Design"
    body = slide.placeholders[1]
    body.text = "Survey design determines the quality of collected data."
    prs.save(str(file_path))

    text = extract_text(file_path, ".pptx")
    assert "Survey Design" in text
    assert "collected data" in text
